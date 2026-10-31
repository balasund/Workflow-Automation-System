"""Gmail-to-workflow.xlsx background polling for Automation Workflow System.

Requires a one-time OAuth setup using authorize_gmail.py. The app reads Gmail
messages only and writes recognized fields into the existing Inprogress sheet.
"""
from __future__ import annotations

import base64
import copy
import json
import os
import re
import shutil
import threading
import time
import uuid
from datetime import datetime
from email.utils import parsedate_to_datetime
from pathlib import Path

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
import requests

ROOT = Path(__file__).resolve().parent
SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]
WORKBOOK_PATH = Path(os.getenv("WORKFLOW_XLSX_PATH", str(ROOT / "workflow.xlsx")))
TOKEN_PATH = Path(os.getenv("GMAIL_TOKEN_PATH", str(ROOT / "token.json")))
PROCESSED_PATH = Path(os.getenv("GMAIL_PROCESSED_IDS_PATH", str(ROOT / "processed_gmail_ids.json")))
NOTIFICATIONS_PATH = Path(os.getenv("GMAIL_NOTIFICATIONS_PATH", str(ROOT / "gmail_job_notifications.json")))
POLL_SECONDS = int(os.getenv("GMAIL_POLL_SECONDS", "60"))
GMAIL_QUERY = os.getenv("GMAIL_JOB_QUERY", 'newer_than:7d -from:me')
_LOCK = threading.Lock()

# Aliases are intentionally conservative; only columns already present in Excel are updated.
ALIASES = {
    "job name": ["job name", "job", "project name", "title", "publication"],
    "job id": ["job id", "job number", "job no", "job#", "job code", "id"],
    "client": ["client", "client name", "customer"],
    "status": ["status", "job status"],
    "receipt time": ["receipt time", "received time", "received date", "email received", "receipt date"],
    "total pages": ["total pages", "pages", "page count", "no of pages", "number of pages"],
    "operator": ["operator", "assigned to", "operator name", "assignee"],
    "complexity": ["complexity", "job complexity"],
    "deadline ist": ["deadline ist", "deadline", "due date", "due time", "deadline date"],
    "notes": ["notes", "remarks", "description", "details"],
    "email subject": ["email subject", "subject"],
    "email from": ["email from", "sender", "from"],
}


def _normal(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _read_processed_ids():
    try:
        data = json.loads(PROCESSED_PATH.read_text(encoding="utf-8"))
        return set(data if isinstance(data, list) else [])
    except (OSError, json.JSONDecodeError):
        return set()


def _save_processed_ids(ids):
    tmp = PROCESSED_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(sorted(ids), indent=2), encoding="utf-8")
    os.replace(tmp, PROCESSED_PATH)


def _load_credentials_from_streamlit_secrets():
    """Load a serialized OAuth token from Streamlit Secrets, if configured."""
    try:
        import streamlit as st
        # Preferred format: [gmail] token_json = '{...the full token.json JSON...}'
        token_json = st.secrets.get("gmail", {}).get("token_json", "")
        # Also accept a top-level key for simpler existing deployments.
        if not token_json:
            token_json = st.secrets.get("GMAIL_TOKEN_JSON", "")
        if not token_json:
            return None
        info = json.loads(str(token_json))
        return Credentials.from_authorized_user_info(info, SCOPES)
    except (ImportError, KeyError, TypeError, ValueError, json.JSONDecodeError):
        return None


def _get_credentials():
    """Use local token.json for development, or Streamlit Secrets in cloud."""
    creds = None
    if TOKEN_PATH.exists():
        try:
            creds = Credentials.from_authorized_user_file(str(TOKEN_PATH), SCOPES)
        except (ValueError, json.JSONDecodeError) as exc:
            raise RuntimeError(
                "Local token.json is not valid OAuth JSON. Generate it with authorize_gmail.py, "
                "or configure [gmail].token_json in Streamlit Secrets."
            ) from exc
    if creds is None:
        creds = _load_credentials_from_streamlit_secrets()
    if creds is None:
        raise FileNotFoundError(
            "Gmail authorization is not configured. For local use, run authorize_gmail.py. "
            "For Streamlit Cloud, add the complete token.json JSON as the value of "
            "[gmail].token_json in App Settings > Secrets. Do not upload token.json to GitHub."
        )
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
        # Do not write refreshed credentials to the source tree in cloud. The refresh token
        # remains in the configured credential source; the access token is refreshed in memory.
        if TOKEN_PATH.exists():
            TOKEN_PATH.write_text(creds.to_json(), encoding="utf-8")
    if not creds.valid:
        raise RuntimeError(
            "Gmail authorization is invalid or expired. Re-authorize locally and update "
            "Streamlit Secrets with the newly generated token.json JSON."
        )
    return creds


def _decode_body(payload):
    """Return readable plain text from Gmail's MIME payload."""
    chunks = []
    def walk(part):
        mime = part.get("mimeType", "")
        body = part.get("body", {})
        data = body.get("data")
        if data and mime == "text/plain":
            try:
                chunks.append(base64.urlsafe_b64decode(data + "==").decode("utf-8", errors="replace"))
            except Exception:
                pass
        for child in part.get("parts", []) or []:
            walk(child)
    walk(payload or {})
    if chunks:
        return "\n".join(chunks)
    # Fallback for HTML-only messages: strip tags and common entities.
    html_chunks = []
    def walk_html(part):
        if part.get("mimeType") == "text/html" and part.get("body", {}).get("data"):
            try:
                html_chunks.append(base64.urlsafe_b64decode(part["body"]["data"] + "==").decode("utf-8", errors="replace"))
            except Exception:
                pass
        for child in part.get("parts", []) or []:
            walk_html(child)
    walk_html(payload or {})
    html = "\n".join(html_chunks)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</li>", "\n", html)
    html = re.sub(r"<[^>]+>", " ", html)
    return html.replace("&nbsp;", " ").replace("&amp;", "&").strip()


def _headers_and_body(service, message_id):
    msg = service.users().messages().get(userId="me", id=message_id, format="full").execute()
    headers = {h.get("name", "").lower(): h.get("value", "") for h in msg.get("payload", {}).get("headers", [])}
    body = _decode_body(msg.get("payload", {}))
    return msg, headers, body


def _parse_labeled_fields(body):
    """Parse common label:value lines; this deliberately avoids guessing from free prose."""
    fields = {}
    for line in body.splitlines():
        line = re.sub(r"^\s*(?:[-*•]|\d+[.)])\s*", "", line).strip()
        match = re.match(r"^([A-Za-z][A-Za-z0-9 _/#().-]{1,45})\s*[:=\-]\s*(.*?)\s*$", line)
        if match and match.group(2):
            fields[_normal(match.group(1))] = match.group(2).strip()
    return fields


def _extract_record(subject, sender, received_at, body):
    labeled = _parse_labeled_fields(body)
    record = {"email subject": subject, "email from": sender}
    for canonical, aliases in ALIASES.items():
        for alias in aliases:
            value = labeled.get(_normal(alias))
            if value:
                record[canonical] = value
                break
    if "job name" not in record:
        # Common subject convention: "New Job: EMLI Brochure".
        m = re.search(r"(?:new\s+job|job\s+request|job\s+name)\s*[:\-]\s*(.+)", subject, re.I)
        if m:
            record["job name"] = m.group(1).strip()
    if "job name" not in record:
        return None, "No clearly labeled Job Name found; skipped for safety."
    record.setdefault("status", "TBI")
    record.setdefault("receipt time", received_at.strftime("%d-%m-%Y %I:%M:%S %p"))
    record.setdefault("notes", body[:1000].strip())
    return record, None


def _find_header_row(ws):
    for row in range(1, min(ws.max_row, 25) + 1):
        vals = [_normal(ws.cell(row, col).value) for col in range(1, ws.max_column + 1)]
        if "status" in vals and any(v in vals for v in ("job name", "job", "project name")):
            return row, vals
    raise ValueError("Could not find a header row containing Job Name and Status in the Inprogress sheet.")


def _column_for_header(header):
    h = _normal(header)
    for canonical, aliases in ALIASES.items():
        if h == _normal(canonical) or h in [_normal(a) for a in aliases]:
            return canonical
    return None


def _github_settings():
    """Read optional GitHub Contents API settings from Streamlit Secrets or environment."""
    try:
        import streamlit as st
        section = st.secrets.get("github", {})
        token = str(section.get("token", "") or st.secrets.get("GITHUB_TOKEN", "")).strip()
        owner = str(section.get("owner", "") or os.getenv("GITHUB_OWNER", "balasund")).strip()
        repo = str(section.get("repo", "") or os.getenv("GITHUB_REPO", "Workflow-Automation-System")).strip()
        branch = str(section.get("branch", "") or os.getenv("GITHUB_BRANCH", "main")).strip()
        path = str(section.get("path", "") or os.getenv("GITHUB_WORKBOOK_PATH", "workflow.xlsx")).strip()
    except Exception:
        token = str(os.getenv("GITHUB_TOKEN", "")).strip()
        owner = str(os.getenv("GITHUB_OWNER", "balasund")).strip()
        repo = str(os.getenv("GITHUB_REPO", "Workflow-Automation-System")).strip()
        branch = str(os.getenv("GITHUB_BRANCH", "main")).strip()
        path = str(os.getenv("GITHUB_WORKBOOK_PATH", "workflow.xlsx")).strip()

    if not token:
        return None
    return {
        "token": token,
        "owner": owner,
        "repo": repo,
        "branch": branch,
        "path": path.lstrip("/"),
    }


def github_sync_configured():
    """Return True only when a GitHub token is configured."""
    return _github_settings() is not None


def _github_api_url(settings):
    from urllib.parse import quote
    encoded_path = "/".join(quote(part, safe="") for part in settings["path"].split("/"))
    return f"https://api.github.com/repos/{settings['owner']}/{settings['repo']}/contents/{encoded_path}"


def _github_headers(settings):
    return {
        "Authorization": f"Bearer {settings['token']}",
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    }


def _download_latest_workbook_from_github():
    """Refresh local workbook from GitHub before editing, preventing stale-file overwrites."""
    settings = _github_settings()
    if not settings:
        return None

    response = requests.get(
        _github_api_url(settings),
        headers=_github_headers(settings),
        params={"ref": settings["branch"]},
        timeout=30,
    )
    if response.status_code == 404:
        raise RuntimeError(
            "GitHub workbook was not found. Check [github].owner, repo, branch, and path in Streamlit Secrets."
        )
    if not response.ok:
        raise RuntimeError(f"Could not download the latest workbook from GitHub (HTTP {response.status_code}): {response.text[:300]}")

    data = response.json()
    if data.get("type") != "file" or not data.get("content"):
        raise RuntimeError("GitHub did not return a downloadable workflow.xlsx file.")
    try:
        workbook_bytes = base64.b64decode(data["content"])
    except Exception as exc:
        raise RuntimeError("GitHub returned invalid base64 content for workflow.xlsx.") from exc

    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    temp_path = WORKBOOK_PATH.with_name(WORKBOOK_PATH.stem + ".github_download_tmp" + WORKBOOK_PATH.suffix)
    temp_path.write_bytes(workbook_bytes)
    os.replace(temp_path, WORKBOOK_PATH)
    return data.get("sha")


def _push_workbook_to_github(expected_sha, job_name):
    """Commit the edited workbook to GitHub using the SHA fetched before editing."""
    settings = _github_settings()
    if not settings:
        return False

    if not expected_sha:
        raise RuntimeError("GitHub did not provide a file SHA; workbook was not uploaded.")

    url = _github_api_url(settings)
    content = base64.b64encode(WORKBOOK_PATH.read_bytes()).decode("ascii")
    payload = {
        "message": f"Update workflow.xlsx from Gmail job: {job_name}",
        "content": content,
        "sha": expected_sha,
        "branch": settings["branch"],
    }
    response = requests.put(url, headers=_github_headers(settings), json=payload, timeout=60)
    if response.status_code in (409, 422):
        raise RuntimeError(
            "GitHub workbook changed during this update. No overwrite was attempted; the next Gmail poll can retry safely."
        )
    if not response.ok:
        raise RuntimeError(f"GitHub upload failed (HTTP {response.status_code}): {response.text[:300]}")
    return True


def _append_record(record):
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Excel workbook not found: {WORKBOOK_PATH}")
    # Workbook writes are serialized within this process. When GitHub sync is configured,
    # always start from the latest remote workbook to preserve other sheets and recent edits.
    with _LOCK:
        github_sha = _download_latest_workbook_from_github()
        wb = load_workbook(WORKBOOK_PATH)
        try:
            sheet_name = next((n for n in wb.sheetnames if _normal(n) == "inprogress"), None)
            if not sheet_name:
                raise ValueError("The workbook has no 'Inprogress' sheet; no changes were made.")
            ws = wb[sheet_name]
            header_row, normalized_headers = _find_header_row(ws)
            col_map = {}
            for col, header in enumerate(normalized_headers, start=1):
                canonical = _column_for_header(header)
                if canonical and canonical not in col_map:
                    col_map[canonical] = col
            if "job name" not in col_map:
                raise ValueError("The Inprogress sheet does not have a recognized Job Name column.")

            job_name_col = col_map["job name"]
            job_id_col = col_map.get("job id")
            wanted_id = str(record.get("job id", "")).strip()
            wanted_name = _normal(record.get("job name"))
            for r in range(header_row + 1, ws.max_row + 1):
                existing_name = _normal(ws.cell(r, job_name_col).value)
                existing_id = str(ws.cell(r, job_id_col).value or "").strip() if job_id_col else ""
                if wanted_id and existing_id == wanted_id:
                    return "duplicate", f"Skipped duplicate job ID {wanted_id}."
                if wanted_name and existing_name == wanted_name:
                    return "duplicate", f"Skipped duplicate job name: {record.get('job name')}."

            target_row = max(ws.max_row + 1, header_row + 1)
            # Avoid writing into a merged range if the workbook has pre-existing merged rows below data.
            while any(not isinstance(ws.cell(target_row, c), MergedCell) for c in range(1, ws.max_column + 1)) is False:
                target_row += 1
            # Copy visual formatting from the previous data row, not its values.
            source_row = target_row - 1
            if source_row > header_row:
                for col in range(1, ws.max_column + 1):
                    src = ws.cell(source_row, col)
                    dst = ws.cell(target_row, col)
                    if isinstance(dst, MergedCell):
                        continue
                    if src.has_style:
                        dst._style = copy.copy(src._style)
                    if src.number_format:
                        dst.number_format = src.number_format
                    if src.alignment:
                        dst.alignment = copy.copy(src.alignment)
                    if src.protection:
                        dst.protection = copy.copy(src.protection)

            written = []
            for canonical, value in record.items():
                col = col_map.get(canonical)
                if not col:
                    continue
                cell = ws.cell(target_row, col)
                if isinstance(cell, MergedCell):
                    continue
                cell.value = value
                written.append(canonical)
            if "job name" not in written:
                raise ValueError("No job name was written; check the Inprogress sheet headers.")

            temp_path = WORKBOOK_PATH.with_name(WORKBOOK_PATH.stem + ".gmail_tmp" + WORKBOOK_PATH.suffix)
            wb.save(temp_path)
            os.replace(temp_path, WORKBOOK_PATH)
            if github_sha:
                _push_workbook_to_github(github_sha, str(record.get("job name", "New job")))
                return "added", f"Added '{record.get('job name')}' to {sheet_name} (row {target_row}) and synced workflow.xlsx to GitHub."
            return "added", f"Added '{record.get('job name')}' to {sheet_name} (row {target_row}). GitHub sync is not configured."
        finally:
            wb.close()


def _write_job_notifications(jobs):
    """Persist notification events for Streamlit to read without a destructive queue."""
    if not jobs:
        return
    with _LOCK:
        try:
            existing = json.loads(NOTIFICATIONS_PATH.read_text(encoding="utf-8"))
            if not isinstance(existing, list):
                existing = []
        except (OSError, json.JSONDecodeError):
            existing = []
        for job in jobs:
            event = dict(job)
            event.setdefault("event_id", uuid.uuid4().hex)
            event.setdefault("created_at", datetime.now().isoformat(timespec="seconds"))
            event.setdefault("job_name", "New job")
            existing.append(event)
        # Keep a bounded history. UI sessions remember which event IDs they have seen.
        existing = existing[-200:]
        tmp = NOTIFICATIONS_PATH.with_suffix(".tmp")
        tmp.write_text(json.dumps(existing, indent=2), encoding="utf-8")
        os.replace(tmp, NOTIFICATIONS_PATH)


def get_pending_job_notifications():
    """Read persisted notification events; do not delete them before the UI consumes them."""
    with _LOCK:
        try:
            events = json.loads(NOTIFICATIONS_PATH.read_text(encoding="utf-8"))
            if not isinstance(events, list):
                return []
        except (OSError, json.JSONDecodeError):
            return []
        # Support events written by earlier versions that had no event_id.
        normalized = []
        changed = False
        for event in events:
            if not isinstance(event, dict):
                continue
            if not event.get("event_id"):
                event = dict(event)
                event["event_id"] = uuid.uuid4().hex
                changed = True
            normalized.append(event)
        if changed:
            tmp = NOTIFICATIONS_PATH.with_suffix(".tmp")
            tmp.write_text(json.dumps(normalized[-200:], indent=2), encoding="utf-8")
            os.replace(tmp, NOTIFICATIONS_PATH)
        return normalized[-200:]


def process_new_emails(max_messages=25, notify_background=False):
    """Poll Gmail, import recognized new job-request emails, and return status messages."""
    service = build("gmail", "v1", credentials=_get_credentials(), cache_discovery=False)
    processed = _read_processed_ids()
    response = service.users().messages().list(userId="me", q=GMAIL_QUERY, maxResults=max_messages).execute()
    messages = response.get("messages", []) or []
    results = []
    for item in reversed(messages):
        msg_id = item["id"]
        if msg_id in processed:
            continue
        try:
            msg, headers, body = _headers_and_body(service, msg_id)
            subject = headers.get("subject", "(no subject)")
            sender = headers.get("from", "")
            internal_ms = int(msg.get("internalDate", "0") or 0)
            received = datetime.fromtimestamp(internal_ms / 1000) if internal_ms else datetime.now()
            record, skip_reason = _extract_record(subject, sender, received, body)
            if skip_reason:
                results.append({"email": subject, "status": "Review required", "detail": skip_reason})
                # Don't mark as processed: it can be imported after correcting the email format.
                continue
            status, detail = _append_record(record)
            results.append({"email": subject, "status": status.title(), "detail": detail})
            if status == "added" and notify_background:
                _write_job_notifications([{
                    "job_name": record.get("job name", "New job"),
                    "detail": detail,
                    "created_at": datetime.now().isoformat(timespec="seconds"),
                }])
            processed.add(msg_id)
            _save_processed_ids(processed)
        except Exception as exc:
            results.append({"email": item.get("id", "Unknown"), "status": "Error", "detail": str(exc)})
    return results


def _poll_forever():
    while True:
        try:
            process_new_emails(notify_background=True)
        except Exception:
            # Streamlit page reports setup/errors; background thread must not kill the app.
            pass
        time.sleep(max(POLL_SECONDS, 30))


def start_background_poller():
    """Start a single daemon poller per Python process."""
    with _LOCK:
        for thread in threading.enumerate():
            if thread.name == "gmail-workflow-poller" and thread.is_alive():
                return False
        thread = threading.Thread(target=_poll_forever, name="gmail-workflow-poller", daemon=True)
        thread.start()
        return True


def gmail_setup_status():
    secrets_token_configured = False
    try:
        import streamlit as st
        secrets_token_configured = bool(
            st.secrets.get("gmail", {}).get("token_json", "")
            or st.secrets.get("GMAIL_TOKEN_JSON", "")
        )
    except Exception:
        pass
    return {
        "credentials_file": (ROOT / "credentials.json").exists(),
        "token_file": TOKEN_PATH.exists() or secrets_token_configured,
        "workbook_file": WORKBOOK_PATH.exists(),
        "poll_seconds": max(POLL_SECONDS, 30),
        "query": GMAIL_QUERY,
        "auth_source": "local token.json" if TOKEN_PATH.exists() else ("Streamlit Secrets" if secrets_token_configured else "not configured"),
        "github_sync_configured": github_sync_configured(),
    }
